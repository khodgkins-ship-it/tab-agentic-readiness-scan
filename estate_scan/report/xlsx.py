"""variants.xlsx -- the working-detail workbook.

One sheet per metric group, every variant with its resolved formula (readable),
usage rank, context, and owner. This is a working artifact: it carries formula
text and owner names, so it travels with the working build, never the
presentation one. Rendered from the (working) findings so it can never disagree
with the report.
"""

import re
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_HEADER = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
_TOP = Alignment(vertical="top")

_COLS = [
    ("Field", 26, "field_name"),
    ("Usage rank", 11, "usage_rank"),
    ("Views", 12, "view_count"),
    ("View share", 11, "view_share"),
    ("Workbooks", 11, "workbook_count"),
    ("Dominant", 10, "is_dominant"),
    ("Resolution", 14, "resolution_status"),
    ("Data source", 24, "datasource_name"),
    ("Owner", 20, "owner"),
    ("Resolved formula", 60, "resolved_formula"),
]

_INVALID = re.compile(r"[\[\]\*\?/\\:]")


def _sheet_title(label, used):
    # type: (str, set) -> str
    base = _INVALID.sub("_", (label or "group")).strip() or "group"
    base = base[:28]
    title = base
    n = 2
    while title.lower() in used:
        suffix = " %d" % n
        title = base[:28 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def write_variants_xlsx(findings, path):
    # type: (dict, str) -> str
    dm = findings.get("findings", {}).get("definition_multiplicity", {})
    groups = dm.get("groups", [])

    wb = Workbook()
    _overview(wb.active, findings, groups)

    used = {"overview"}
    for g in groups:
        ws = wb.create_sheet(_sheet_title(g.get("label"), used))
        _group_sheet(ws, g)

    wb.save(path)
    return path


def _overview(ws, findings, groups):
    ws.title = "Overview"
    meta = findings.get("meta", {})
    ws["A1"] = "Metric variants"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Run %s  ·  build %s  ·  generated %s" % (
        meta.get("run_id"), meta.get("build"), meta.get("generated_at"))
    ws["A3"] = ("Working detail: contains resolved formulas and owner names. "
                "Do not circulate with the presentation build.")
    ws["A3"].font = Font(italic=True)

    head = ["Metric", "Variants", "Workbooks", "Disagreeing", "Cover 80%",
            "Dominant"]
    row = 5
    for i, h in enumerate(head, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = _HEADER
    for g in groups:
        row += 1
        ws.cell(row=row, column=1, value=g.get("label"))
        ws.cell(row=row, column=2, value=g.get("variant_count"))
        ws.cell(row=row, column=3, value=g.get("workbooks_affected"))
        ws.cell(row=row, column=4, value=g.get("disagreeing_variants"))
        ws.cell(row=row, column=5, value=g.get("variants_covering_80pct_views"))
        ws.cell(row=row, column=6,
                value="yes" if g.get("dominant") else "no")
    widths = [30, 10, 11, 12, 10, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"


def _group_sheet(ws, g):
    ws["A1"] = g.get("label")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "%d variants  ·  %s" % (
        g.get("variant_count") or 0,
        "dominant variant present" if g.get("dominant") else "no dominant variant")

    header_row = 4
    for i, (label, width, _key) in enumerate(_COLS, start=1):
        c = ws.cell(row=header_row, column=i, value=label)
        c.font = _HEADER
        ws.column_dimensions[get_column_letter(i)].width = width

    r = header_row
    for v in g.get("variants", []):
        r += 1
        for i, (_label, _w, key) in enumerate(_COLS, start=1):
            val = v.get(key)
            if key == "is_dominant":
                val = "yes" if val else ""
            elif key == "view_share" and isinstance(val, (int, float)):
                val = round(val, 4)
            cell = ws.cell(row=r, column=i, value=val)
            cell.alignment = _WRAP if key == "resolved_formula" else _TOP
    ws.freeze_panes = "A%d" % (header_row + 1)
